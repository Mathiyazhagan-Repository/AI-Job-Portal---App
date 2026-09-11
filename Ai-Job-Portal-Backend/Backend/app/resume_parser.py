from __future__ import annotations

import io
import os
import re
from datetime import datetime
from typing import Any

from fastapi import APIRouter, HTTPException, UploadFile, status
from pypdf import PdfReader
from docx import Document

router = APIRouter(prefix="/api")
MAX_UPLOAD_BYTES = int(os.getenv("MAX_UPLOAD_BYTES", "5242880"))

EMAIL_RE = re.compile(r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}", re.I)
PHONE_RE = re.compile(
    r"(?:\+?\d{1,3}[-.\s]?)?(?:\(?\d{2,4}\)?[-.\s]?)\d{3}[-.\s]?\d{4}",
    re.I,
)
LINK_RE = re.compile(r"https?://[^\s\)\]>\"]+", re.I)
HEADING_RE = re.compile(r"^(?:summary|professional summary|profile|objective|skills|education|experience|work experience|projects|certifications|achievements|languages|additional information|other relevant information)\b.*$", re.I)

KNOWN_SKILLS = [
    "python", "javascript", "typescript", "react", "node", "node.js", "next.js", "java",
    "sql", "postgresql", "mysql", "mongodb", "redis", "aws", "azure", "gcp", "docker",
    "kubernetes", "terraform", "html", "css", "tailwind", "redux", "graphql", "jest",
    "cypress", "playwright", "git", "linux", "fastapi", "flask", "django", "rest api",
    "microservices", "data structures", "algorithm", "machine learning", "ai", "vue", "angular",
    "figma", "design systems", "excel", "power bi", "tableau", "spark", "pandas", "numpy",
    "pytorch", "tensorflow", "c++", "c#", ".net", "rust", "go", "ruby", "php", "laravel",
    "solidity", "blockchain", "ui/ux", "product management", "agile", "scrum",
]

KNOWN_LOCATIONS = [
    "bengaluru", "bangalore", "hyderabad", "chennai", "pune", "delhi", "mumbai",
    "kolkata", "ahmedabad", "jaipur", "kochi", "gurugram", "gurgaon", "noida",
    "new delhi", "san francisco", "new york", "london", "singapore", "dubai",
]

SECTION_ALIASES = {
    "summary": ["summary", "professional summary", "profile", "about", "overview"],
    "skills": ["skills", "technical skills", "core competencies", "tools", "technologies"],
    "education": ["education", "academic background", "education background", "academics"],
    "experience": ["experience", "work experience", "professional experience", "employment history"],
    "projects": ["projects", "project work", "selected projects"],
    "certifications": ["certifications", "licenses", "training", "certificates"],
    "achievements": ["achievements", "awards", "honors", "accomplishments"],
    "languages": ["languages", "language proficiency"],
}


def normalize_spaces(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def normalize_phone(value: str | None) -> str | None:
    if not value:
        return None
    digits = re.sub(r"\D+", "", value)
    if len(digits) in (10,):
        return f"+91 {digits[:5]} {digits[5:]}" if digits.startswith("9") else digits
    if len(digits) >= 10:
        return value.strip()
    return value.strip() if value.strip() else None


def normalize_date(value: str | None) -> str | None:
    if not value:
        return None
    cleaned = normalize_spaces(value).strip("., ")
    if not cleaned:
        return None
    return cleaned


def first_match(regex: re.Pattern[str], text: str) -> str | None:
    match = regex.search(text)
    if not match:
        return None
    return match.group(0).strip()


def infer_name(text: str) -> str | None:
    lines = [normalize_spaces(line) for line in text.splitlines() if normalize_spaces(line)]
    for line in lines[:12]:
        candidate = line
        if '@' in candidate:
            candidate = candidate.split('@', 1)[0].strip()
        candidate = re.sub(r"https?://\S+", "", candidate)
        candidate = re.sub(r"[|•\-–—]", " ", candidate)
        candidate = re.sub(r"\s+", " ", candidate).strip()
        if not candidate or re.search(r"\d", candidate):
            continue
        if re.search(r"(summary|education|experience|skills|projects|certifications|linkedin|github|portfolio|achievements|languages)", candidate, re.I):
            continue
        if re.fullmatch(r"[A-Za-z][A-Za-z0-9._-]*", candidate) and " " not in candidate:
            continue
        words = re.findall(r"[A-Za-z][A-Za-z.'-]+", candidate)
        if len(words) >= 2:
            name = " ".join(words[:4])
            if not re.fullmatch(r"(?:[A-Za-z]+[.'-]?\s+){1,3}[A-Za-z]+[.'-]?", name):
                continue
            return name.title()
    return None


def extract_location(text: str) -> str | None:
    for loc in KNOWN_LOCATIONS:
        if re.search(rf"\b{re.escape(loc)}\b", text, re.I):
            match = re.search(rf"\b{re.escape(loc)}\b.*?(?:,|\||\n|$)", text, re.I)
            if match:
                return normalize_spaces(match.group(0).rstrip(",|\n"))
    return None


def extract_links(text: str) -> dict[str, str | None]:
    urls = LINK_RE.findall(text)
    result = {"linkedin": None, "github": None, "portfolio": None}
    for url in urls:
        u = url.lower()
        if "linkedin.com" in u:
            result["linkedin"] = url.strip()
        elif "github.com" in u:
            result["github"] = url.strip()
        elif result["portfolio"] is None:
            result["portfolio"] = url.strip()
    return result


def extract_summary(text: str) -> str | None:
    patterns = [
        r"(?:professional\s+summary|summary|profile|about|overview)\s*[:\-]?\s*(.*?)(?=(?:\n\s*(?:skills|education|experience|work experience|projects|certifications|achievements|languages|technical skills)\b)|$)",
    ]
    for pattern in patterns:
        match = re.search(pattern, text, re.I | re.S)
        if match:
            value = normalize_spaces(match.group(1))
            if len(value) > 10:
                return value
    return None


def extract_skills(text: str) -> list[str]:
    found: set[str] = set()
    skill_names = {skill.lower(): skill for skill in KNOWN_SKILLS}

    sections = split_sections(text)
    block = section_text(
        sections,
        ["skills", "technical skills", "core competencies", "tools", "tools & platforms", "technology stack"],
    )

    if block:
        for chunk in re.split(r"[\n|;•]+", block):
            for part in re.split(r",", chunk):
                item = normalize_spaces(part).strip(".,;|•-")
                if not item or len(item) > 40:
                    continue
                lower = item.lower()
                if lower in {"skills", "tools", "platforms", "technical skills", "core competencies"}:
                    continue
                if re.search(r"\b(?:education|experience|projects|achievements|certifications|languages|summary|profile|objective)\b", item, re.I):
                    continue
                if lower in skill_names:
                    found.add(skill_names[lower])
                elif any(lower == key or lower.startswith(f"{key} ") or lower.endswith(f" {key}") for key in skill_names):
                    for key, label in skill_names.items():
                        if lower == key or lower.startswith(f"{key} ") or lower.endswith(f" {key}"):
                            found.add(label)
                            break

    if not found:
        for skill in KNOWN_SKILLS:
            if re.search(rf"(?<![A-Za-z]){re.escape(skill.lower())}(?![A-Za-z])", text.lower()):
                found.add(skill)

    return sorted(found, key=lambda x: x.lower())


def split_sections(text: str) -> list[tuple[str, str]]:
    normalized = text.replace("\r", "\n")
    lines = normalized.split("\n")
    sections: list[tuple[str, str]] = []
    current_heading = ""
    current_body: list[str] = []
    section_aliases = (
        "summary", "profile", "objective", "skills", "education", "experience",
        "work experience", "projects", "certifications", "achievements", "languages",
        "technical skills", "core competencies", "additional information", "other relevant information",
        "tools", "tools & platforms", "technology stack",
    )

    def flush() -> None:
        if current_heading:
            sections.append((current_heading.strip().lower(), "\n".join(current_body).strip()))

    for line in lines:
        cleaned = line.strip()
        if not cleaned:
            if current_heading:
                current_body.append("")
            continue

        heading_match = re.match(r"^(?:[\*\-•\s]*)([A-Za-z][A-Za-z /&()\-:]{1,80})\s*(?::|-)?\s*(.*)$", cleaned)
        if heading_match:
            heading = heading_match.group(1).strip().lower()
            remainder = heading_match.group(2).strip()
            is_heading = heading in section_aliases or any(alias in heading for alias in ("summary", "skills", "education", "experience", "projects", "certifications", "achievements", "languages", "tools"))
            if is_heading:
                flush()
                current_heading = heading_match.group(1).strip()
                current_body = []
                if remainder:
                    current_body.append(remainder)
                continue
        if current_heading:
            current_body.append(cleaned)
    flush()
    return sections


def section_text(sections: list[tuple[str, str]], names: list[str]) -> str:
    for name, content in sections:
        if name.lower() in names:
            return content
    for name, content in sections:
        if any(alias in name.lower() for alias in names):
            return content
    return ""


def parse_education_entries(text: str) -> list[dict[str, Any]]:
    sections = split_sections(text)
    education_block = section_text(sections, ["education", "academic background", "education background", "academics"])
    if not education_block:
        education_match = re.search(
            r"(?is)(?:education|academic background)\s*[:\-]?\s*(.*?)(?=(?:\n\s*(?:skills|experience|projects|certifications|achievements|languages|summary|tools)\b)|$)",
            text,
        )
        if education_match:
            education_block = education_match.group(1)

    if not education_block:
        return []

    entries: list[dict[str, Any]] = []
    blocks = re.split(r"\n\s*(?:\n\s*)+(?=[A-Z][A-Za-z0-9 .,'&()-]+\s*(?:\n|$))", education_block)
    for block in blocks:
        block_text = normalize_spaces(block)
        if not block_text:
            continue

        institution_match = re.search(r"([A-Z][A-Za-z0-9 .&'-]+(?:University|College|Institute|School|Academy|Institute of Technology))", block_text)
        institution = institution_match.group(1).strip() if institution_match else None
        degree_match = re.search(r"\b(B\.Tech|BTech|B\.E\.|BE|B\.Sc|BS|M\.Tech|MTech|MBA|B\.A|B\.Com|M\.A|MS|PhD|Diploma|Bachelor|Master|Certificate)\b", block_text, re.I)
        degree = degree_match.group(0) if degree_match else None
        field_match = re.search(r"(?:in|of)\s+([A-Za-z][A-Za-z &/.-]{3,80})", block_text, re.I)
        field = field_match.group(1).strip() if field_match else None
        date_match = re.search(r"(?:20\d{2}|19\d{2})\s*(?:-|–|to|\s*)\s*(?:Present|20\d{2}|19\d{2})", block_text, re.I)
        dates = date_match.group(0) if date_match else None
        grade_match = re.search(r"(?:GPA|CGPA|Grade|Percentage|Score)\s*[:\-]?\s*([0-9.]+\s*(?:/\s*[0-9.]+|%|CGPA|GPA)?)", block_text, re.I)
        grade = grade_match.group(1) if grade_match else None

        if institution or degree:
            entries.append({
                "institution": institution,
                "degree": degree,
                "field_of_study": field,
                "start_date": None,
                "end_date": None,
                "grade": grade,
            })
    return entries


def parse_experience_entries(text: str) -> list[dict[str, Any]]:
    sections = split_sections(text)
    exp_block = section_text(sections, ["experience", "work experience", "professional experience", "employment history"])
    if not exp_block:
        return []

    entries: list[dict[str, Any]] = []
    split_candidates = re.split(r"\n\s*(?=(?:[A-Z][A-Za-z0-9 .&'-]+)\s*(?:\n|$))", exp_block)
    for chunk in split_candidates:
        item = normalize_spaces(chunk)
        if len(item) < 30:
            continue
        company_match = re.search(r"([A-Z][A-Za-z0-9 .&'-]+(?:Inc|Ltd|Private|LLP|Pvt|Labs|Technologies|Systems|Solutions|Health|Consulting))", item)
        title_match = re.search(r"(?:\b(?:Senior|Junior|Lead|Engineer|Developer|Manager|Analyst|Consultant|Specialist|Associate|Designer|Architect)\b[ A-Za-z0-9\-./&]*)", item)
        date_match = re.search(r"(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec|20\d{2}|19\d{2})[A-Za-z0-9 .,-]{0,25}(?:to|–|-|Present)", item, re.I)

        start_date = None
        end_date = None
        if date_match:
            date_text = date_match.group(0).replace("–", "-").replace("—", "-")
            date_text = re.sub(r"\s+", " ", date_text).strip()
            parts = re.split(r"\s*(?:to|-)\s*", date_text, maxsplit=1)
            if len(parts) >= 2:
                start_date = normalize_date(parts[0].strip())
                end_date = normalize_date(parts[1].strip())
            else:
                start_date = normalize_date(date_text)
                end_date = "Present" if re.search(r"\bpresent\b", date_text, re.I) else None

        entries.append({
            "company": company_match.group(1).strip() if company_match else None,
            "job_title": title_match.group(0).strip() if title_match else None,
            "start_date": start_date,
            "end_date": end_date,
            "description": item[:500] if item else None,
        })
    return entries


def parse_projects(text: str) -> list[dict[str, Any]]:
    sections = split_sections(text)
    project_block = section_text(sections, ["projects", "project work", "selected projects"])
    if not project_block:
        return []

    projects: list[dict[str, Any]] = []
    chunks = re.split(r"\n\s*(?=(?:[A-Z][A-Za-z0-9 &()/-]+\s*(?:\n|$)))", project_block)
    for chunk in chunks:
        item = normalize_spaces(chunk)
        if len(item) < 30:
            continue
        url = first_match(LINK_RE, item)
        description = item[:400]
        tech = [s.strip() for s in re.split(r"[,•\-\n]", item) if len(s.strip()) > 1][:8]
        projects.append({
            "project_name": item.split(" ")[:6] and item.split(" ")[:6]
            and " ".join(item.split(" ")[:6]) or None,
            "description": description,
            "technologies_used": tech,
            "project_url": url,
        })
    return projects


def parse_list_section(text: str, names: list[str]) -> list[str]:
    sections = split_sections(text)
    block = section_text(sections, names)
    if not block:
        return []
    items: list[str] = []
    for line in re.split(r"\n|•|\-|;", block):
        cleaned = normalize_spaces(line)
        if cleaned and len(cleaned) > 1:
            items.append(cleaned)
    return items


def _parse_exp_date(date_text: str | None) -> datetime | None:
    if not date_text:
        return None

    value = normalize_spaces(date_text)
    if not value:
        return None

    value = value.replace("–", "-").replace("—", "-")
    value = re.sub(r"\s+", " ", value)

    if value.lower() == "present":
        return datetime.now()

    if "present" in value.lower():
        value = re.sub(r"\s*-\s*present\b", "", value, flags=re.I)
        value = value.strip()
        if not value:
            return datetime.now()

    for fmt in ("%b %Y", "%B %Y", "%Y", "%b-%Y", "%B-%Y", "%Y-%m", "%m/%Y", "%m-%Y", "%d %b %Y", "%d %B %Y"):
        try:
            return datetime.strptime(value, fmt)
        except ValueError:
            continue

    year_match = re.search(r"(19|20)\d{2}", value)
    if year_match:
        try:
            return datetime.strptime(year_match.group(1) + year_match.group(0)[2:], "%Y")
        except ValueError:
            return datetime.strptime(year_match.group(0), "%Y")

    return None


def calculate_total_experience_years(experience_entries: list[dict[str, Any]]) -> float | None:
    if not experience_entries:
        return None

    total_months = 0.0
    valid = False
    for item in experience_entries:
        start = _parse_exp_date(item.get("start_date"))
        end = _parse_exp_date(item.get("end_date"))
        if start is None:
            continue
        end_value = end or datetime.now()
        months = (end_value.year - start.year) * 12 + (end_value.month - start.month)
        if months < 0:
            continue
        total_months += months
        valid = True

    if not valid:
        return None

    return round(total_months / 12, 1)


def extract_current_title(text: str, experience_entries: list[dict[str, Any]] | None = None) -> str | None:
    if experience_entries:
        for item in experience_entries:
            if item.get("job_title"):
                return item.get("job_title")

    exp_block = section_text(split_sections(text), ["experience", "work experience", "professional experience", "employment history"])
    if exp_block:
        lines = [normalize_spaces(line) for line in exp_block.splitlines() if normalize_spaces(line)]
        for line in lines[:6]:
            if not re.search(r"(summary|education|skills|projects|certifications|languages|achievements)", line, re.I):
                if not re.search(r"\b(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec|20\d{2}|19\d{2})\b", line, re.I):
                    return line
    return None


def parse_resume_text(text: str) -> dict[str, Any]:
    cleaned = re.sub(r"\u00a0+", " ", text)
    cleaned = cleaned.replace("\r\n", "\n").replace("\r", "\n")
    cleaned = cleaned.strip()

    email = first_match(EMAIL_RE, cleaned)
    phone = normalize_phone(first_match(PHONE_RE, cleaned))
    location = extract_location(cleaned)
    links = extract_links(cleaned)
    linkedin = links.get("linkedin")
    github = links.get("github")
    portfolio = links.get("portfolio")
    name = infer_name(cleaned)

    summary = extract_summary(cleaned)
    skills = extract_skills(cleaned)
    education = parse_education_entries(cleaned)
    experience = parse_experience_entries(cleaned)
    current_title = extract_current_title(cleaned, experience)
    total_experience_years = calculate_total_experience_years(experience)
    projects = parse_projects(cleaned)
    certifications = parse_list_section(cleaned, ["certifications", "licenses", "training", "certificates"])
    achievements = parse_list_section(cleaned, ["achievements", "awards", "honors", "accomplishments"])
    languages = parse_list_section(cleaned, ["languages", "language proficiency"])

    other_sections: list[dict[str, Any]] = []
    for heading, body in split_sections(cleaned):
        heading_text = heading.strip()
        if heading_text.lower() in {"summary", "profile", "skills", "education", "experience", "work experience", "projects", "certifications", "achievements", "languages"}:
            continue
        if body and len(body) > 10:
            other_sections.append({"title": heading_text, "content": body})

    return {
        "personal_info": {
            "name": name,
            "email": email,
            "phone": phone,
            "location": location,
            "linkedin": linkedin,
            "github": github,
            "portfolio": portfolio,
        },
        "summary": summary,
        "current_title": current_title,
        "total_experience_years": total_experience_years,
        "skills": skills,
        "education": education,
        "experience": experience,
        "projects": projects,
        "certifications": certifications,
        "achievements": achievements,
        "languages": languages,
        "other_sections": other_sections,
    }


def extract_text_from_pdf(file_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(file_bytes))
    pages: list[str] = []
    for page in reader.pages:
        extracted = page.extract_text() or ""
        pages.append(extracted)
    return "\n".join(pages)


def extract_text_from_docx(file_bytes: bytes) -> str:
    doc = Document(io.BytesIO(file_bytes))
    paragraphs: list[str] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            paragraphs.append(text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                value = cell.text.strip()
                if value:
                    paragraphs.append(value)
    return "\n".join(paragraphs)


def parse_resume_file_bytes(file_name: str, file_bytes: bytes) -> dict[str, Any]:
    if not file_name:
        raise HTTPException(status_code=400, detail="Resume file name is required.")
    lower_name = file_name.lower()
    if lower_name.endswith(".pdf"):
        text = extract_text_from_pdf(file_bytes)
    elif lower_name.endswith((".doc", ".docx")):
        text = extract_text_from_docx(file_bytes)
    else:
        raise HTTPException(status_code=415, detail="Unsupported file type. Upload a PDF or DOCX resume.")

    if not text or not text.strip():
        raise HTTPException(status_code=422, detail="The uploaded resume is empty or no text could be extracted.")

    result = parse_resume_text(text)
    if not any(result.get("personal_info", {}).values()) and not result.get("skills") and not result.get("experience") and not result.get("education"):
        raise HTTPException(status_code=422, detail="Text could not be extracted from the resume.")
    return result


@router.post("/resume/parse")
async def parse_resume_endpoint(file: UploadFile):
    if not file or not file.filename:
        raise HTTPException(status_code=400, detail="No file was uploaded.")

    file_name = file.filename.lower()
    if not (file_name.endswith(".pdf") or file_name.endswith(".doc") or file_name.endswith(".docx")):
        raise HTTPException(status_code=415, detail="Unsupported file type. Only PDF and DOCX files are supported.")

    content = await file.read()
    if len(content) == 0:
        raise HTTPException(status_code=422, detail="The uploaded resume is empty.")
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="The file is too large. Maximum supported size is 5 MB.")

    try:
        return parse_resume_file_bytes(file.filename, content)
    except HTTPException:
        raise
    except Exception as exc:  # pragma: no cover - defensive backend guard
        raise HTTPException(status_code=422, detail=f"Parsing failed: {str(exc)}") from exc
